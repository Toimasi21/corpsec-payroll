import os
import pytest
import io
import zipfile
import openpyxl
from application.backend.app.models.audit import AuditLog
from application.backend.app.models.payroll import PayrollPeriod, PayrollRecord, PayrollError


def get_admin_headers(client):
    login_resp = client.post(
        "/api/v1/auth/login",
        json={"email": "admin@corpsec.co.ke", "password": os.getenv("ADMIN_INITIAL_PASSWORD", "TestAdminSecretPass123!")}
    )
    token = login_resp.json()["access_token"]
    return {"Authorization": f"Bearer {token}"}


@pytest.fixture
def calculated_period(client):
    headers = get_admin_headers(client)
    
    # Create guard with missing statutory NSSF number to generate a validation error
    client.post(
        "/api/v1/guards",
        json={
            "employee_number": "CS-00888",
            "full_name": "Incomplete Guard",
            "national_id": "88776655",
            "phone": "0788776655",
            "hire_date": "2024-01-01",
            "basic_salary": 15000.0,
            "payment_method": "Bank Transfer",
            "nssf_number": "",  # Empty NSSF -> triggers error
            "shif_number": "SHIF-888"
        },
        headers=headers
    )

    # Create period
    p_resp = client.post(
        "/api/v1/payroll/periods",
        json={
            "year": 2026,
            "month": 9,
            "period_name": "September 2026",
            "start_date": "2026-09-01",
            "end_date": "2026-09-30"
        },
        headers=headers
    )
    period_id = p_resp.json()["id"]

    # Calculate
    calc_resp = client.post(f"/api/v1/payroll/{period_id}/calculate", headers=headers)
    assert calc_resp.status_code == 200

    return period_id


def test_error_ignore_workflow(client, db, calculated_period):
    headers = get_admin_headers(client)
    
    # Get errors for period
    errors_resp = client.get(f"/api/v1/payroll/{calculated_period}/errors", headers=headers)
    errors = errors_resp.json()
    assert len(errors) >= 1
    target_err = errors[0]

    # Ignore error with reason
    resolve_resp = client.post(
        f"/api/v1/payroll/errors/{target_err['id']}/resolve",
        json={"status": "IGNORED", "reason": "Guard NSSF number application in progress at ministry"},
        headers=headers
    )
    assert resolve_resp.status_code == 200
    assert resolve_resp.json()["status"] == "IGNORED"

    # Verify audit log entry
    db.expire_all()
    audit = db.query(AuditLog).filter(AuditLog.action == "IGNORE_PAYROLL_ERROR").first()
    assert audit is not None
    assert "application in progress" in audit.reason


def test_payroll_confirmation(client, db, calculated_period):
    headers = get_admin_headers(client)

    # First resolve any open errors if present
    errors_resp = client.get(f"/api/v1/payroll/{calculated_period}/errors", headers=headers)
    for err in errors_resp.json():
        if err["status"] == "OPEN":
            client.post(
                f"/api/v1/payroll/errors/{err['id']}/resolve",
                json={"status": "IGNORED", "reason": "Resolved for testing"},
                headers=headers
            )

    # Confirm payroll
    confirm_resp = client.post(f"/api/v1/payroll/{calculated_period}/confirm", headers=headers)
    assert confirm_resp.status_code == 200
    assert confirm_resp.json()["status"] == "CONFIRMED"

    # Verify period in DB
    db.expire_all()
    period = db.query(PayrollPeriod).filter(PayrollPeriod.id == calculated_period).first()
    assert period.status == "CONFIRMED"


def test_payslip_preview_and_pdf(client, calculated_period):
    headers = get_admin_headers(client)
    
    records = client.get(f"/api/v1/payroll/{calculated_period}/records", headers=headers).json()
    assert len(records) >= 1
    rec_id = records[0]["id"]

    # Preview
    preview_resp = client.get(f"/api/v1/payslips/preview/{rec_id}", headers=headers)
    assert preview_resp.status_code == 200
    data = preview_resp.json()
    assert "gross_pay" in data
    assert "net_pay" in data

    # Download PDF
    pdf_resp = client.get(f"/api/v1/payslips/{rec_id}/pdf", headers=headers)
    assert pdf_resp.status_code == 200
    assert pdf_resp.headers["content-type"] == "application/pdf"
    assert pdf_resp.content.startswith(b"%PDF")


def test_bulk_payslips_zip_export(client, calculated_period):
    headers = get_admin_headers(client)

    response = client.post(f"/api/v1/payslips/bulk-generate?period_id={calculated_period}", headers=headers)
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/zip"

    # Validate ZIP archive
    zip_buf = io.BytesIO(response.content)
    with zipfile.ZipFile(zip_buf, "r") as zf:
        file_list = zf.namelist()
        assert len(file_list) >= 1
        assert any(f.endswith(".pdf") for f in file_list)


def test_bank_payment_schedule_excel(client, calculated_period):
    headers = get_admin_headers(client)

    # Fetch JSON Bank Schedule
    json_resp = client.get(f"/api/v1/payments/bank/{calculated_period}", headers=headers)
    assert json_resp.status_code == 200
    bdata = json_resp.json()
    assert "total_amount" in bdata

    # Export Excel
    excel_resp = client.get(f"/api/v1/payments/bank/{calculated_period}/export-excel", headers=headers)
    assert excel_resp.status_code == 200
    assert "spreadsheetml.sheet" in excel_resp.headers["content-type"]

    # Validate Excel workbook
    wb = openpyxl.load_workbook(io.BytesIO(excel_resp.content))
    assert "Bank Payment Schedule" in wb.sheetnames
    ws = wb["Bank Payment Schedule"]
    assert ws["A1"].value == "CORPSEC PAYROLL — BANK PAYMENT SCHEDULE"


def test_mpesa_payment_schedule_excel(client, calculated_period):
    headers = get_admin_headers(client)

    # Fetch JSON M-Pesa Schedule
    json_resp = client.get(f"/api/v1/payments/mpesa/{calculated_period}", headers=headers)
    assert json_resp.status_code == 200

    # Export Excel
    excel_resp = client.get(f"/api/v1/payments/mpesa/{calculated_period}/export-excel", headers=headers)
    assert excel_resp.status_code == 200
    assert "spreadsheetml.sheet" in excel_resp.headers["content-type"]

    wb = openpyxl.load_workbook(io.BytesIO(excel_resp.content))
    assert "M-Pesa Payment Schedule" in wb.sheetnames
