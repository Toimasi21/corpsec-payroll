import os
import pytest
import io
from datetime import date
import openpyxl
from application.backend.app.models.guard import Guard
from application.backend.app.models.payroll import PayrollPeriod, PayrollRecord
from application.backend.app.models.user import User
from application.backend.app.models.audit import AuditLog
from application.backend.app.services.payment_service import (
    validate_kenyan_phone,
    validate_bank_details,
    mask_sensitive_account,
    build_payment_schedules_and_reconciliation,
)


def get_admin_headers(client):
    login_resp = client.post(
        "/api/v1/auth/login",
        json={"email": "admin@corpsec.co.ke", "password": os.getenv("ADMIN_INITIAL_PASSWORD", "TestAdminSecretPass123!")}
    )
    token = login_resp.json()["access_token"]
    return {"Authorization": f"Bearer {token}"}


def get_guard_headers(client):
    login_resp = client.post(
        "/api/v1/auth/login/guard",
        json={"guard_id": "CS-00452", "password": "Guard@123456"}
    )
    token = login_resp.json()["access_token"]
    return {"Authorization": f"Bearer {token}"}


@pytest.fixture
def setup_payment_test_data(db):
    """
    Creates a dedicated payroll period with diverse guards:
    - Guard 1: Bank Transfer (Valid Bank Details)
    - Guard 2: M-Pesa (Valid 07... Phone)
    - Guard 3: Bank Transfer (Missing Bank Account -> Unrouted)
    - Guard 4: M-Pesa (Invalid Phone -> Unrouted)
    - Guard 5: Bank Transfer (Duplicate Account with Guard 1)
    """
    period = PayrollPeriod(
        year=2026,
        month=10,
        period_name="October 2026",
        start_date=date(2026, 10, 1),
        end_date=date(2026, 10, 31),
        status="CONFIRMED"
    )
    db.add(period)
    db.commit()
    db.refresh(period)

    # 1. Valid Bank Guard
    g1 = Guard(
        employee_number="CS-90001",
        full_name="Alice Wanjiku",
        national_id="90000001",
        phone="0711000001",
        hire_date=date(2024, 1, 1),
        payment_method="Bank Transfer",
        bank_name="KCB Bank",
        bank_account="1122334455",
        status="ACTIVE"
    )
    # 2. Valid M-Pesa Guard
    g2 = Guard(
        employee_number="CS-90002",
        full_name="Bob Ochieng",
        national_id="90000002",
        phone="0722000002",
        mpesa_number="0722000002",
        hire_date=date(2024, 1, 1),
        payment_method="M-Pesa",
        status="ACTIVE"
    )
    # 3. Invalid Bank Guard (missing account)
    g3 = Guard(
        employee_number="CS-90003",
        full_name="Charlie Mutua",
        national_id="90000003",
        phone="0733000003",
        hire_date=date(2024, 1, 1),
        payment_method="Bank Transfer",
        bank_name="Equity Bank",
        bank_account="000000000",  # Dummy invalid account
        status="ACTIVE"
    )
    # 4. Invalid M-Pesa Guard (malformed phone)
    g4 = Guard(
        employee_number="CS-90004",
        full_name="David Kiprop",
        national_id="90000004",
        phone="INVALID123",
        mpesa_number="INVALID123",
        hire_date=date(2024, 1, 1),
        payment_method="M-Pesa",
        status="ACTIVE"
    )
    # 5. Duplicate Bank Account Guard
    g5 = Guard(
        employee_number="CS-90005",
        full_name="Eve Adhiambo",
        national_id="90000005",
        phone="0755000005",
        hire_date=date(2024, 1, 1),
        payment_method="Bank Transfer",
        bank_name="KCB Bank",
        bank_account="1122334455",  # Duplicate of g1
        status="ACTIVE"
    )

    db.add_all([g1, g2, g3, g4, g5])
    db.commit()
    db.refresh(g1)
    db.refresh(g2)
    db.refresh(g3)
    db.refresh(g4)
    db.refresh(g5)

    # Payroll records
    r1 = PayrollRecord(payroll_period_id=period.id, guard_id=g1.id, days_worked=20, regular_hours=160.0, overtime_hours=0.0, gross_pay=25000.0, net_pay=20000.0)
    r2 = PayrollRecord(payroll_period_id=period.id, guard_id=g2.id, days_worked=20, regular_hours=160.0, overtime_hours=0.0, gross_pay=22000.0, net_pay=18000.0)
    r3 = PayrollRecord(payroll_period_id=period.id, guard_id=g3.id, days_worked=20, regular_hours=160.0, overtime_hours=0.0, gross_pay=20000.0, net_pay=16000.0)
    r4 = PayrollRecord(payroll_period_id=period.id, guard_id=g4.id, days_worked=20, regular_hours=160.0, overtime_hours=0.0, gross_pay=21000.0, net_pay=17000.0)
    r5 = PayrollRecord(payroll_period_id=period.id, guard_id=g5.id, days_worked=20, regular_hours=160.0, overtime_hours=0.0, gross_pay=24000.0, net_pay=19000.0)

    db.add_all([r1, r2, r3, r4, r5])
    db.commit()

    return period.id


def test_phone_and_bank_unit_validators():
    # Valid Kenyan Phones
    val, phone, err = validate_kenyan_phone("0712345678")
    assert val is True
    assert phone == "0712345678"

    val, phone, err = validate_kenyan_phone("+254712345678")
    assert val is True
    assert phone == "0712345678"

    val, phone, err = validate_kenyan_phone("0112345678")
    assert val is True
    assert phone == "0112345678"

    # Invalid Phones
    val, _, err = validate_kenyan_phone("12345")
    assert val is False

    val, _, err = validate_kenyan_phone(None)
    assert val is False

    # Valid Bank
    val, err = validate_bank_details("KCB Bank", "123456789")
    assert val is True

    # Invalid Bank
    val, err = validate_bank_details("N/A", "123456789")
    assert val is False

    val, err = validate_bank_details("Equity Bank", "000000000")
    assert val is False

    # Account Masking
    masked = mask_sensitive_account("1122334455")
    assert masked == "******4455"
    assert "112233" not in masked


def test_payment_routing_and_reconciliation(client, db, setup_payment_test_data):
    period_id = setup_payment_test_data
    admin_headers = get_admin_headers(client)

    resp = client.get(f"/api/v1/payments/reconciliation/{period_id}", headers=admin_headers)
    assert resp.status_code == 200
    data = resp.json()

    assert data["period_id"] == period_id
    assert data["is_reconciled"] is True
    assert data["discrepancy"] == 0.0

    # Total Net = 20000 + 18000 + 16000 + 17000 + 19000 = 90,000.0
    assert data["total_payroll_net_pay"] == 90000.0
    # Bank = 20000 (g1) + 19000 (g5) = 39,000.0
    assert data["total_bank_amount"] == 39000.0
    # M-Pesa = 18000 (g2) = 18,000.0
    assert data["total_mpesa_amount"] == 18000.0
    # Unrouted = 16000 (g3) + 17000 (g4) = 33,000.0
    assert data["total_unrouted_amount"] == 33000.0

    # Duplicate Warnings
    assert len(data["duplicate_warnings"]) >= 1
    dup = data["duplicate_warnings"][0]
    assert dup["type"] == "DUPLICATE_BANK_ACCOUNT"

    # Verify Sum(Bank) + Sum(Mpesa) + Sum(Unrouted) == Total Net Pay
    assert round(data["total_bank_amount"] + data["total_mpesa_amount"] + data["total_unrouted_amount"], 2) == data["total_payroll_net_pay"]


def test_bank_schedule_routing_and_exclusion(client, setup_payment_test_data):
    period_id = setup_payment_test_data
    admin_headers = get_admin_headers(client)

    resp = client.get(f"/api/v1/payments/bank/{period_id}", headers=admin_headers)
    assert resp.status_code == 200
    data = resp.json()

    # Should contain CS-90001 and CS-90005, but NOT CS-90002 (M-Pesa) or CS-90003 (invalid account)
    emp_numbers = [item["employee_number"] for item in data["items"]]
    assert "CS-90001" in emp_numbers
    assert "CS-90005" in emp_numbers
    assert "CS-90002" not in emp_numbers
    assert "CS-90003" not in emp_numbers

    assert data["total_amount"] == 39000.0


def test_mpesa_schedule_routing_and_exclusion(client, setup_payment_test_data):
    period_id = setup_payment_test_data
    admin_headers = get_admin_headers(client)

    resp = client.get(f"/api/v1/payments/mpesa/{period_id}", headers=admin_headers)
    assert resp.status_code == 200
    data = resp.json()

    # Should contain CS-90002, but NOT CS-90001 (Bank) or CS-90004 (invalid phone)
    emp_numbers = [item["employee_number"] for item in data["items"]]
    assert "CS-90002" in emp_numbers
    assert "CS-90001" not in emp_numbers
    assert "CS-90004" not in emp_numbers

    assert data["total_amount"] == 18000.0


def test_cross_period_isolation(client, db, setup_payment_test_data):
    period_id_1 = setup_payment_test_data
    admin_headers = get_admin_headers(client)

    # Create second empty period
    p2 = PayrollPeriod(
        year=2026,
        month=11,
        period_name="November 2026",
        start_date=date(2026, 11, 1),
        end_date=date(2026, 11, 30),
        status="DRAFT"
    )
    db.add(p2)
    db.commit()
    db.refresh(p2)

    resp = client.get(f"/api/v1/payments/bank/{p2.id}", headers=admin_headers)
    assert resp.status_code == 200
    data = resp.json()
    assert len(data["items"]) == 0
    assert data["total_amount"] == 0.0


def test_excel_export_bank_format_and_openpyxl(client, setup_payment_test_data):
    period_id = setup_payment_test_data
    admin_headers = get_admin_headers(client)

    resp = client.get(f"/api/v1/payments/bank/{period_id}/export-excel", headers=admin_headers)
    assert resp.status_code == 200
    assert "spreadsheetml.sheet" in resp.headers["content-type"]

    wb = openpyxl.load_workbook(io.BytesIO(resp.content))
    assert "Bank Payment Schedule" in wb.sheetnames
    ws = wb["Bank Payment Schedule"]

    # Header Row Titles check (Row 4)
    expected_headers = ["Guard ID", "Full Name", "Bank Name", "Account Number", "Net Amount (KES)", "Reference", "Status"]
    actual_headers = [ws.cell(row=4, column=c).value for c in range(1, 8)]
    assert actual_headers == expected_headers

    # Freeze panes check
    assert ws.freeze_panes == "A5"

    # Numeric formatting check
    data_cell = ws.cell(row=5, column=5)
    assert data_cell.number_format == "#,##0.00"
    assert isinstance(data_cell.value, (int, float))

    # Total row check
    total_label = ws.cell(row=7, column=1).value
    assert total_label == "TOTAL"


def test_excel_export_mpesa_format_and_openpyxl(client, setup_payment_test_data):
    period_id = setup_payment_test_data
    admin_headers = get_admin_headers(client)

    resp = client.get(f"/api/v1/payments/mpesa/{period_id}/export-excel", headers=admin_headers)
    assert resp.status_code == 200

    wb = openpyxl.load_workbook(io.BytesIO(resp.content))
    ws = wb["M-Pesa Payment Schedule"]

    # Header Row Titles check (Row 4)
    expected_headers = ["Guard ID", "Full Name", "Phone Number", "Net Amount (KES)", "Reference", "Status"]
    actual_headers = [ws.cell(row=4, column=c).value for c in range(1, 7)]
    assert actual_headers == expected_headers

    # Freeze panes check
    assert ws.freeze_panes == "A5"


def test_payment_security_and_rbac(client, setup_payment_test_data):
    period_id = setup_payment_test_data

    # Guard token should be forbidden (403)
    guard_headers = get_guard_headers(client)
    res_guard = client.get(f"/api/v1/payments/bank/{period_id}", headers=guard_headers)
    assert res_guard.status_code == 403
    assert "Administrative privileges required" in res_guard.json()["detail"]

    # Unauthenticated should be unauthorized (401)
    res_unauth = client.get(f"/api/v1/payments/bank/{period_id}")
    assert res_unauth.status_code == 401
